import re
import fitz  # PyMuPDF
from docx import Document
from typing import Dict, List

# Large skill keywords list
SKILLS_LIST = [
    # Programming Languages
    "python", "java", "javascript", "typescript", "c++", "c#", "c", "ruby", "go", "rust",
    "kotlin", "swift", "php", "scala", "r", "matlab", "perl", "dart", "html", "css",
    # Web Frameworks
    "react", "angular", "vue", "nextjs", "nodejs", "express", "django", "flask", "fastapi",
    "spring", "laravel", "rails", "asp.net", "tailwind", "bootstrap",
    # AI/ML
    "machine learning", "deep learning", "natural language processing", "nlp", "computer vision",
    "tensorflow", "pytorch", "keras", "scikit-learn", "pandas", "numpy", "matplotlib",
    "seaborn", "opencv", "huggingface", "transformers", "bert", "gpt", "llm",
    "reinforcement learning", "neural network", "cnn", "rnn", "lstm", "gan",
    # Data
    "sql", "mysql", "postgresql", "mongodb", "sqlite", "redis", "elasticsearch",
    "data analysis", "data science", "data visualization", "tableau", "power bi",
    "hadoop", "spark", "kafka", "airflow", "etl",
    # Cloud & DevOps
    "aws", "azure", "google cloud", "gcp", "docker", "kubernetes", "jenkins",
    "ci/cd", "git", "github", "gitlab", "linux", "bash", "terraform", "ansible",
    # Mobile
    "android", "ios", "react native", "flutter", "xamarin",
    # Tools & Others
    "rest api", "graphql", "microservices", "agile", "scrum", "jira", "postman",
    "selenium", "junit", "pytest", "unit testing", "api testing",
    "blockchain", "iot", "cybersecurity", "networking", "cloud computing",
]


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from PDF using PyMuPDF"""
    text = ""
    try:
        doc = fitz.open(file_path)
        for page in doc:
            text += page.get_text()
        doc.close()
    except Exception as e:
        print(f"PDF extraction error: {e}")
    return text


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from DOCX"""
    text = ""
    try:
        doc = Document(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
        # Also read tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
            text += "\n"
    except Exception as e:
        print(f"DOCX extraction error: {e}")
    return text


def extract_text(file_path: str) -> str:
    """Auto-detect file type and extract text"""
    if file_path.lower().endswith(".pdf"):
        return extract_text_from_pdf(file_path)
    elif file_path.lower().endswith(".docx"):
        return extract_text_from_docx(file_path)
    return ""


def extract_email(text: str) -> str:
    """Extract email address from text"""
    pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
    matches = re.findall(pattern, text)
    return matches[0] if matches else ""


def extract_phone(text: str) -> str:
    """Extract phone number from text"""
    pattern = r'[\+]?[(]?[0-9]{3}[)]?[-\s\.]?[0-9]{3}[-\s\.]?[0-9]{4,6}'
    matches = re.findall(pattern, text)
    return matches[0] if matches else ""


def extract_name(text: str) -> str:
    """Extract candidate name (first non-empty line usually)"""
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    for line in lines[:5]:
        # Skip lines that look like emails, phones, or URLs
        if not re.search(r'[@\d/\\|]', line) and len(line.split()) <= 5 and len(line) > 3:
            return line
    return "Unknown"


def extract_skills(text: str) -> List[str]:
    """Extract skills by matching against known skills list"""
    text_lower = text.lower()
    found_skills = []
    for skill in SKILLS_LIST:
        if skill.lower() in text_lower:
            found_skills.append(skill)
    return list(set(found_skills))


def extract_experience_years(text: str) -> int:
    """Extract years of experience from text"""
    patterns = [
        r'(\d+)\+?\s*years?\s*of\s*experience',
        r'(\d+)\+?\s*years?\s*experience',
        r'experience\s*of\s*(\d+)\+?\s*years?',
        r'(\d+)\+?\s*yr[s]?\s*exp',
    ]
    for pattern in patterns:
        match = re.search(pattern, text.lower())
        if match:
            return int(match.group(1))
    return 0


def extract_education(text: str) -> str:
    """Extract highest education level"""
    text_lower = text.lower()
    if any(x in text_lower for x in ["ph.d", "phd", "doctorate"]):
        return "PhD"
    elif any(x in text_lower for x in ["m.tech", "m.e.", "mtech", "master of technology"]):
        return "M.Tech"
    elif any(x in text_lower for x in ["mba", "master of business"]):
        return "MBA"
    elif any(x in text_lower for x in ["m.sc", "msc", "master of science"]):
        return "M.Sc"
    elif any(x in text_lower for x in ["b.tech", "b.e.", "btech", "bachelor of technology", "bachelor of engineering"]):
        return "B.Tech"
    elif any(x in text_lower for x in ["bca", "bachelor of computer"]):
        return "BCA"
    elif any(x in text_lower for x in ["b.sc", "bsc", "bachelor of science"]):
        return "B.Sc"
    elif any(x in text_lower for x in ["12th", "hsc", "higher secondary", "intermediate"]):
        return "12th"
    elif any(x in text_lower for x in ["10th", "ssc", "matriculation"]):
        return "10th"
    return "Not Specified"


def parse_resume(file_path: str) -> Dict:
    """Full resume parser — returns structured data"""
    text = extract_text(file_path)
    if not text:
        return {}

    return {
        "raw_text": text,
        "name": extract_name(text),
        "email": extract_email(text),
        "phone": extract_phone(text),
        "skills": extract_skills(text),
        "experience_years": extract_experience_years(text),
        "education": extract_education(text),
    }
